"""Generate HTML, PDF, and DOCX for Laporan Pekerjaan Jasa Pengembangan Software.
Target: Exactly 3 pages matching the PDF reference.
"""

import os
import subprocess
import base64
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.abspath("Docs")
IMAGES_DIR = os.path.join(DOCS_DIR, "images")

def encode_image(filename):
    path = os.path.join(IMAGES_DIR, filename)
    if os.path.exists(path):
        with open(path, "rb") as f:
            return "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')
    return ""

img1_b64 = encode_image("gambar1_arsitektur_software.png")
img2_b64 = encode_image("gambar2_kebutuhan_modul.png")
img3_b64 = encode_image("gambar3_gui_radar.png")
img4_b64 = encode_image("gambar4_integrasi_pengujian.png")

html_content = f"""<!DOCTYPE html>
<html lang="id">
<head>
<meta charset="UTF-8">
<title>Laporan Pekerjaan - Jasa Pengembangan Software</title>
<style>
    @page {{
        size: A4 portrait;
        margin: 16mm 20mm 16mm 20mm;
    }}
    * {{
        box-sizing: border-box;
    }}
    body {{
        font-family: Arial, "Helvetica Neue", Helvetica, sans-serif;
        font-size: 10.2pt;
        line-height: 1.34;
        color: #000000;
        margin: 0;
        padding: 0;
    }}
    .page {{
        page-break-after: always;
        break-after: page;
    }}
    .page:last-child {{
        page-break-after: avoid;
        break-after: avoid;
    }}
    .header-title {{
        text-align: center;
        margin-bottom: 16px;
    }}
    .header-title h1 {{
        font-size: 14pt;
        font-weight: bold;
        margin: 0 0 4px 0;
    }}
    .header-title h2 {{
        font-size: 12.5pt;
        font-weight: bold;
        margin: 0;
    }}
    h3 {{
        font-size: 10.5pt;
        font-weight: bold;
        margin: 12px 0 5px 0;
    }}
    p {{
        text-align: justify;
        text-justify: inter-word;
        margin: 0 0 8px 0;
        line-height: 1.35;
    }}
    .figure-container {{
        text-align: center;
        margin: 6px 0 4px 0;
    }}
    .figure-container img {{
        max-width: 96%;
        height: auto;
        border: 1px solid #cbd5e1;
        border-radius: 3px;
    }}
    .caption {{
        font-size: 9.5pt;
        text-align: center;
        margin: 4px 0 8px 0;
        color: #0f172a;
    }}
    /* Signature Block */
    .signature-table {{
        width: 100%;
        border-collapse: collapse;
        margin-top: 14px;
        font-size: 9.5pt;
    }}
    .signature-table td {{
        border: 1px solid #000000;
        vertical-align: top;
        text-align: center;
        padding: 0;
    }}
    .sig-header {{
        padding: 5px;
        font-weight: bold;
    }}
    .sig-header-orange {{
        background-color: #f59e0b;
        color: #000000;
        padding: 5px;
        font-weight: bold;
    }}
    .sig-space {{
        height: 52px;
        display: flex;
        align-items: center;
        justify-content: center;
        color: #9ca3af;
        font-style: italic;
        font-size: 8.5pt;
    }}
    .sig-space-orange {{
        background-color: #f59e0b;
        height: 52px;
        display: flex;
        align-items: center;
        justify-content: center;
        color: #111827;
        font-style: italic;
        font-weight: 500;
        font-size: 8.5pt;
    }}
    .sig-footer {{
        border-top: 1px solid #000000;
        padding: 5px 4px;
        line-height: 1.22;
        font-size: 9pt;
    }}
    .sig-footer-orange {{
        background-color: #f59e0b;
        border-top: 1px solid #000000;
        padding: 5px 4px;
        line-height: 1.22;
        font-size: 9pt;
        font-weight: 500;
    }}
</style>
</head>
<body>

<!-- ================= PAGE 1 ================= -->
<div class="page">
    <div class="header-title">
        <h1>Laporan Pekerjaan</h1>
        <h2>Jasa Pengembangan Software</h2>
    </div>

    <h3>1. Pendahuluan</h3>
    <p>
        Pada penelitian ini, diusulkan dan dikembangkan sistem radar FMCW (Frequency Modulated Continuous Wave) 
        berbasis perangkat lunak terintegrasi yang mampu melakukan pemantauan, deteksi jarak, dan estimasi posisi 
        target secara presisi dan waktu-nyata (real-time). Sistem ini dirancang untuk mendukung penelitian strategis 
        dalam skema LPDP RISPRO bekerja sama dengan Institut Teknologi Bandung (DKST ITB). Sistem ini mengintegrasikan 
        perangkat keras akuisisi data berkecepatan tinggi (high-speed ADC) dengan sistem mekanik pemindai antena 
        otomatis guna meminimalkan intervensi manual dan memastikan keandalan deteksi gelombang mikro. Bagian 
        terpenting pada sistem ini adalah perangkat lunak kendali dan pengolahan sinyal digital yang bertugas 
        mengakuisisi data gelombang IF (Intermediate Frequency), melakukan analisis spektrum frekuensi beat, 
        serta memvisualisasikan target pada layar radar.
    </p>

    <h3>2. Pengembangan Software Radar</h3>
    <div class="figure-container">
        <img src="{img1_b64}" alt="Gambar 1. Desain Arsitektur Sistem Software Radar" style="max-height: 76mm;">
    </div>
    <div class="caption">Gambar 1. Desain Arsitektur Sistem Software Radar</div>

    <p>
        Gambar 1 menampilkan desain arsitektur keseluruhan dari sistem perangkat lunak radar yang dikembangkan. 
        Desain ini dibuat dalam struktur modular berbasis multi-threading untuk menjamin aliran data yang kontinu 
        tanpa hambatan (non-blocking). Komponen utama dalam arsitektur ini adalah integrasi kartu ADC ADLink 
        PCI-9846H dengan resolusi 16-bit dan kecepatan sampling 20 MS/s yang dihubungkan melalui pustaka C 
        (wd-dask64.dll) menggunakan antarmuka ctypes pada Python. Sistem akuisisi dikonfigurasi dengan mode 
        asynchronous continuous restart double-buffer DMA dan trigger digital eksternal (Ext-D), memungkinkan 
        pemindahan data langsung ke memori RAM (zero-copy memory mapping) tanpa membebani media penyimpanan. 
        Secara paralel, worker serial membaca sudut antena dari mikrokontroler penggerak motor secara kontinu, 
        kemudian seluruh data dialirkan melalui thread-safe queue menuju inti pemrosesan sinyal digital (DSP) 
        dan antarmuka visualisasi Dear PyGui.
    </p>
</div>

<!-- ================= PAGE 2 ================= -->
<div class="page">
    <div class="figure-container" style="margin-top: 2px;">
        <img src="{img2_b64}" alt="Gambar 2. Kebutuhan modul dan komponen sistem software radar" style="max-height: 72mm;">
    </div>
    <div class="caption">Gambar 2. Kebutuhan modul dan komponen sistem software radar</div>

    <p>
        Gambar 2 merupakan daftar kebutuhan modul perangkat lunak (software bill of deliverables) yang dirancang 
        dan diimplementasikan dalam pengembangan sistem radar ini. Daftar ini mencakup berbagai komponen inti seperti 
        modul C-DAQ engine untuk akuisisi DMA, modul streaming memori zero-copy, pustaka pemroses sinyal FFT 
        berbasis scipy, algoritma penghalus spektrum Savitzky-Golay orde 3, algoritma deteksi puncak target, 
        widget radar PPI taktis, widget penganalisis spektrum daya RF (dBm pada beban 50 Ohm), osiloskop domain 
        waktu untuk 20.000 sampel per buffer, modul sinkronisasi serial antena pemindai, hingga tema antarmuka 
        taktis kedirgantaraan (aerospace dark palette). Semua modul dipilih dan dibangun berdasarkan keandalan, 
        kecepatan komputasi numerik, dan kompatibilitas sistem. Keseluruhan modul telah diuji dan diintegrasikan 
        secara komprehensif sehingga siap dioperasikan tanpa ada kendala teknis dalam proses instalasi dan pengoperasian.
    </p>

    <div class="figure-container" style="margin-top: 4px;">
        <img src="{img3_b64}" alt="Gambar 3. Antarmuka Pengguna (GUI) Sistem Pemantauan Radar Real-Time" style="max-height: 82mm;">
    </div>
    <div class="caption">Gambar 3. Antarmuka Pengguna (GUI) Sistem Pemantauan Radar Real-Time</div>
</div>

<!-- ================= PAGE 3 ================= -->
<div class="page">
    <div class="figure-container" style="margin-top: 2px;">
        <img src="{img4_b64}" alt="Gambar 4. Pengujian Integrasi dan Validasi Sistem Software Radar" style="max-height: 60mm;">
    </div>
    <div class="caption">Gambar 4. Pengujian Integrasi dan Validasi Sistem Software Radar</div>

    <p>
        Pekerjaan jasa pengembangan software sistem radar telah selesai dilakukan sesuai dengan spesifikasi teknis 
        yang telah ditentukan. Proses pengembangan dimulai dari penataan arsitektur modular, pembuatan pustaka 
        antarmuka driver C untuk komunikasi perangkat keras akuisisi data, serta perancangan tata letak antarmuka 
        grafis pengguna (GUI) berbasis Dear PyGui. Tahap awal perakitan perangkat lunak dilakukan dengan menyusun 
        pipeline akuisisi data multi-kanal (CH0 dan CH2) berkecepatan 20 MS/s menggunakan teknik asynchronous 
        DMA continuous restart. Integrasi ini memastikan transfer data gelombang mikro dari ADC langsung diterima 
        oleh memori Python secara real-time. Modul pemrosesan sinyal digital kemudian dipasang untuk melakukan 
        penghilangan offset DC, penerapan windowing Hann dengan normalisasi gain koheren, perhitungan Fast Fourier 
        Transform (FFT), serta kalibrasi daya fisik RF dalam satuan dBm pada impedansi 50 Ohm.
    </p>

    <p>
        Pada Gambar 3 seluruh subsistem antarmuka ditampilkan dalam kondisi beroperasi penuh (live execution). 
        Sistem visualisasi radar PPI menampilkan sektor pemindaian 0° hingga 180° dengan cincin jarak 0 sampai 15 km, 
        dilengkapi jarum pemindai bergradasi phosphor green yang bergerak sinkron dengan sudut aktual antena. 
        Penganalisis spektrum frekuensi beat berhasil meredam derau rumput (noise grass) menggunakan filter 
        Savitzky-Golay, memungkinkan deteksi puncak sinyal target secara akurat. Osiloskop domain waktu menampilkan 
        gelombang sinusoidal intermediate frequency (IF) secara simultan dengan tabel telemetri yang menyajikan 
        informasi jarak, frekuensi puncak, dan estimasi daya RF secara real-time. Bilah status ribbon di bagian atas 
        menampilkan telemetri sistem yang mencakup kestabilan streaming pada 60 FPS, status trigger, dan pencatat event.
    </p>

    <p>
        Pada Gambar 4 seluruh tahapan pengujian dan validasi sistem dilakukan secara menyeluruh guna menjamin 
        stabilitas dan akurasi perangkat lunak. Pengujian meliputi verifikasi jabat-tangan driver ADC PCI-9846H, 
        sinkronisasi serial komunikasi motor penggerak antena pada kecepatan 115.200 bps, validasi streaming data 
        tanpa packet drop pada throughput 40 MB/detik per kanal, uji performa algoritma FFT dan filter derau, 
        hingga pengujian ketahanan memori jangka panjang (zero memory leak) menggunakan circular buffer deque O(1). 
        Semua parameter fungsional telah diperiksa dan memenuhi kriteria yang telah ditetapkan. Dengan selesainya 
        seluruh rangkaian kegiatan tersebut, pekerjaan jasa pengembangan software ini dinyatakan selesai.
    </p>

    <!-- Signature Table matching PDF Page 3 -->
    <table class="signature-table">
        <tr>
            <td style="width: 50%;">
                <div class="sig-header">Menyetujui<br>Ketua Peneliti</div>
                <div class="sig-space">TTD Peneliti</div>
                <div class="sig-footer">
                    <strong>Dr. Ir. Eko Mursito Budi, M.T.</strong><br>
                    NIP : 196710061997021001
                </div>
            </td>
            <td style="width: 50%;">
                <div class="sig-header-orange">Perusahaan / Pelaksana<br>CV. Maja Baru Indonesia</div>
                <div class="sig-space-orange">TTD Perusahaan</div>
                <div class="sig-footer-orange">
                    <strong>Aris Munandar</strong><br>
                    Direktur
                </div>
            </td>
        </tr>
    </table>
</div>

</body>
</html>
"""

# Write HTML file
html_path = os.path.join(DOCS_DIR, "Laporan_Jasa_Pengembangan_Software.html")
with open(html_path, "w", encoding="utf-8") as f:
    f.write(html_content)
print(f"Generated HTML: {html_path}")

# Convert HTML to PDF using Chrome Headless
pdf_path = os.path.join(DOCS_DIR, "Laporan_Jasa_Pengembangan_Software.pdf")
chrome_exe = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
if os.path.exists(chrome_exe):
    cmd = [
        chrome_exe,
        "--headless",
        "--disable-gpu",
        "--run-all-compositor-stages-before-draw",
        "--print-to-pdf-no-header",
        f"--print-to-pdf={pdf_path}",
        html_path
    ]
    try:
        subprocess.run(cmd, check=True)
        print(f"Generated PDF via Chrome: {pdf_path}")
        
        with open(pdf_path, 'rb') as f:
            pdf_data = f.read()
        page_matches = re.findall(rb'/Type\s*/Page\b', pdf_data)
        print(f"Verified PDF page count: {len(page_matches)}")
    except Exception as e:
        print(f"Chrome PDF generation error: {e}")

# ================= DOCX GENERATION =================
docx_path = os.path.join(DOCS_DIR, "Laporan_Jasa_Pengembangan_Software.docx")
doc = Document()

# Set page margins to match PDF
for section in doc.sections:
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.page_width = Inches(8.27)  # A4
    section.page_height = Inches(11.69)

# Normal style font
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(10)
font.color.rgb = RGBColor(0, 0, 0)

def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.size = Pt(9.5)
    run.font.name = 'Arial'

def add_heading_section(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.name = 'Arial'

# Page 1
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_title.paragraph_format.space_before = Pt(0)
p_title.paragraph_format.space_after = Pt(2)
run1 = p_title.add_run("Laporan Pekerjaan\n")
run1.font.size = Pt(14)
run1.font.bold = True
run1.font.name = 'Arial'

run2 = p_title.add_run("Jasa Pengembangan Software")
run2.font.size = Pt(12.5)
run2.font.bold = True
run2.font.name = 'Arial'

add_heading_section("1. Pendahuluan")
p1 = doc.add_paragraph(
    "Pada penelitian ini, diusulkan dan dikembangkan sistem radar FMCW (Frequency Modulated Continuous Wave) "
    "berbasis perangkat lunak terintegrasi yang mampu melakukan pemantauan, deteksi jarak, dan estimasi posisi "
    "target secara presisi dan waktu-nyata (real-time). Sistem ini dirancang untuk mendukung penelitian strategis "
    "dalam skema LPDP RISPRO bekerja sama dengan Institut Teknologi Bandung (DKST ITB). Sistem ini mengintegrasikan "
    "perangkat keras akuisisi data berkecepatan tinggi (high-speed ADC) dengan sistem mekanik pemindai antena "
    "otomatis guna meminimalkan intervensi manual dan memastikan keandalan deteksi gelombang mikro. Bagian "
    "terpenting pada sistem ini adalah perangkat lunak kendali dan pengolahan sinyal digital yang bertugas "
    "mengakuisisi data gelombang IF (Intermediate Frequency), melakukan analisis spektrum frekuensi beat, "
    "serta memvisualisasikan target pada layar radar."
)
p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p1.paragraph_format.space_after = Pt(6)
p1.paragraph_format.line_spacing = 1.25

add_heading_section("2. Pengembangan Software Radar")
img1_file = os.path.join(IMAGES_DIR, "gambar1_arsitektur_software.png")
if os.path.exists(img1_file):
    doc.add_picture(img1_file, width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Gambar 1. Desain Arsitektur Sistem Software Radar")

p2 = doc.add_paragraph(
    "Gambar 1 menampilkan desain arsitektur keseluruhan dari sistem perangkat lunak radar yang dikembangkan. "
    "Desain ini dibuat dalam struktur modular berbasis multi-threading untuk menjamin aliran data yang kontinu "
    "tanpa hambatan (non-blocking). Komponen utama dalam arsitektur ini adalah integrasi kartu ADC ADLink "
    "PCI-9846H dengan resolusi 16-bit dan kecepatan sampling 20 MS/s yang dihubungkan melalui pustaka C "
    "(wd-dask64.dll) menggunakan antarmuka ctypes pada Python. Sistem akuisisi dikonfigurasi dengan mode "
    "asynchronous continuous restart double-buffer DMA dan trigger digital eksternal (Ext-D), memungkinkan "
    "pemindahan data langsung ke memori RAM (zero-copy memory mapping) tanpa membebani media penyimpanan. "
    "Secara paralel, worker serial membaca sudut antena dari mikrokontroler penggerak motor secara kontinu, "
    "kemudian seluruh data dialirkan melalui thread-safe queue menuju inti pemrosesan sinyal digital (DSP) "
    "dan antarmuka visualisasi Dear PyGui."
)
p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p2.paragraph_format.space_after = Pt(6)
p2.paragraph_format.line_spacing = 1.25

# Page 2
doc.add_page_break()

img2_file = os.path.join(IMAGES_DIR, "gambar2_kebutuhan_modul.png")
if os.path.exists(img2_file):
    doc.add_picture(img2_file, width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Gambar 2. Kebutuhan modul dan komponen sistem software radar")

p3 = doc.add_paragraph(
    "Gambar 2 merupakan daftar kebutuhan modul perangkat lunak (software bill of deliverables) yang dirancang "
    "dan diimplementasikan dalam pengembangan sistem radar ini. Daftar ini mencakup berbagai komponen inti seperti "
    "modul C-DAQ engine untuk akuisisi DMA, modul streaming memori zero-copy, pustaka pemroses sinyal FFT "
    "berbasis scipy, algoritma penghalus spektrum Savitzky-Golay orde 3, algoritma deteksi puncak target, "
    "widget radar PPI taktis, widget penganalisis spektrum daya RF (dBm pada beban 50 Ohm), osiloskop domain "
    "waktu untuk 20.000 sampel per buffer, modul sinkronisasi serial antena pemindai, hingga tema antarmuka "
    "taktis kedirgantaraan (aerospace dark palette). Semua modul dipilih dan dibangun berdasarkan keandalan, "
    "kecepatan komputasi numerik, dan kompatibilitas sistem. Keseluruhan modul telah diuji dan diintegrasikan "
    "secara komprehensif sehingga siap dioperasikan tanpa ada kendala teknis dalam proses instalasi dan pengoperasian."
)
p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p3.paragraph_format.space_after = Pt(6)
p3.paragraph_format.line_spacing = 1.25

img3_file = os.path.join(IMAGES_DIR, "gambar3_gui_radar.png")
if os.path.exists(img3_file):
    doc.add_picture(img3_file, width=Inches(5.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Gambar 3. Antarmuka Pengguna (GUI) Sistem Pemantauan Radar Real-Time")

# Page 3
doc.add_page_break()

img4_file = os.path.join(IMAGES_DIR, "gambar4_integrasi_pengujian.png")
if os.path.exists(img4_file):
    doc.add_picture(img4_file, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
add_caption("Gambar 4. Pengujian Integrasi dan Validasi Sistem Software Radar")

p4 = doc.add_paragraph(
    "Pekerjaan jasa pengembangan software sistem radar telah selesai dilakukan sesuai dengan spesifikasi teknis "
    "yang telah ditentukan. Proses pengembangan dimulai dari penataan arsitektur modular, pembuatan pustaka "
    "antarmuka driver C untuk komunikasi perangkat keras akuisisi data, serta perancangan tata letak antarmuka "
    "grafis pengguna (GUI) berbasis Dear PyGui. Tahap awal perakitan perangkat lunak dilakukan dengan menyusun "
    "pipeline akuisisi data multi-kanal (CH0 dan CH2) berkecepatan 20 MS/s menggunakan teknik asynchronous "
    "DMA continuous restart. Integrasi ini memastikan transfer data gelombang mikro dari ADC langsung diterima "
    "oleh memori Python secara real-time. Modul pemrosesan sinyal digital kemudian dipasang untuk melakukan "
    "penghilangan offset DC, penerapan windowing Hann dengan normalisasi gain koheren, perhitungan Fast Fourier "
    "Transform (FFT), serta kalibrasi daya fisik RF dalam satuan dBm pada impedansi 50 Ohm."
)
p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p4.paragraph_format.space_after = Pt(5)
p4.paragraph_format.line_spacing = 1.22

p5 = doc.add_paragraph(
    "Pada Gambar 3 seluruh subsistem antarmuka ditampilkan dalam kondisi beroperasi penuh (live execution). "
    "Sistem visualisasi radar PPI menampilkan sektor pemindaian 0° hingga 180° dengan cincin jarak 0 sampai 15 km, "
    "dilengkapi jarum pemindai bergradasi phosphor green yang bergerak sinkron dengan sudut aktual antena. "
    "Penganalisis spektrum frekuensi beat berhasil meredam derau rumput (noise grass) menggunakan filter "
    "Savitzky-Golay, memungkinkan deteksi puncak sinyal target secara akurat. Osiloskop domain waktu menampilkan "
    "gelombang sinusoidal intermediate frequency (IF) secara simultan dengan tabel telemetri yang menyajikan "
    "informasi jarak, frekuensi puncak, dan estimasi daya RF secara real-time. Bilah status ribbon di bagian atas "
    "menampilkan telemetri sistem yang mencakup kestabilan streaming pada 60 FPS, status trigger, dan pencatat event."
)
p5.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p5.paragraph_format.space_after = Pt(5)
p5.paragraph_format.line_spacing = 1.22

p6 = doc.add_paragraph(
    "Pada Gambar 4 seluruh tahapan pengujian dan validasi sistem dilakukan secara menyeluruh guna menjamin "
    "stabilitas dan akurasi perangkat lunak. Pengujian meliputi verifikasi jabat-tangan driver ADC PCI-9846H, "
    "sinkronisasi serial komunikasi motor penggerak antena pada kecepatan 115.200 bps, validasi streaming data "
    "tanpa packet drop pada throughput 40 MB/detik per kanal, uji performa algoritma FFT dan filter derau, "
    "hingga pengujian ketahanan memori jangka panjang (zero memory leak) menggunakan circular buffer deque O(1). "
    "Semua parameter fungsional telah diperiksa dan memenuhi kriteria yang telah ditetapkan. Dengan selesainya "
    "seluruh rangkaian kegiatan tersebut, pekerjaan jasa pengembangan software ini dinyatakan selesai."
)
p6.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p6.paragraph_format.space_after = Pt(8)
p6.paragraph_format.line_spacing = 1.22

# Signature Table matching PDF
table = doc.add_table(rows=3, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
col_widths = [Inches(3.2), Inches(3.2)]

# Row 0: Headers
cell_00 = table.cell(0, 0)
cell_00.text = "Menyetujui\nKetua Peneliti"
cell_00.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
cell_00.paragraphs[0].runs[0].font.bold = True
cell_00.paragraphs[0].runs[0].font.size = Pt(9.5)

cell_01 = table.cell(0, 1)
cell_01.text = "Perusahaan / Pelaksana\nCV. Maja Baru Indonesia"
cell_01.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
cell_01.paragraphs[0].runs[0].font.bold = True
cell_01.paragraphs[0].runs[0].font.size = Pt(9.5)

# Row 1: Space for signatures
cell_10 = table.cell(1, 0)
cell_10.text = "\n\nTTD Peneliti\n\n"
cell_10.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
cell_10.paragraphs[0].runs[0].font.italic = True
cell_10.paragraphs[0].runs[0].font.size = Pt(8.5)
cell_10.paragraphs[0].runs[0].font.color.rgb = RGBColor(140, 140, 140)

cell_11 = table.cell(1, 1)
cell_11.text = "\n\nTTD Perusahaan\n\n"
cell_11.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
cell_11.paragraphs[0].runs[0].font.italic = True
cell_11.paragraphs[0].runs[0].font.size = Pt(8.5)
cell_11.paragraphs[0].runs[0].font.color.rgb = RGBColor(140, 140, 140)

# Row 2: Footers / Names
cell_20 = table.cell(2, 0)
cell_20.text = "Dr. Ir. Eko Mursito Budi, M.T.\nNIP : 196710061997021001"
cell_20.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
cell_20.paragraphs[0].runs[0].font.bold = True
cell_20.paragraphs[0].runs[0].font.size = Pt(9)

cell_21 = table.cell(2, 1)
cell_21.text = "Aris Munandar\nDirektur"
cell_21.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
cell_21.paragraphs[0].runs[0].font.bold = True
cell_21.paragraphs[0].runs[0].font.size = Pt(9)

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

for r_idx, row in enumerate(table.rows):
    for c_idx, cell in enumerate(row.cells):
        cell.width = col_widths[c_idx]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_border(cell, top=dict(sz=6, val='single', color='000000'),
                              bottom=dict(sz=6, val='single', color='000000'),
                              left=dict(sz=6, val='single', color='000000'),
                              right=dict(sz=6, val='single', color='000000'))
        if c_idx == 1:
            set_cell_shading(cell, "F59E0B")

doc.save(docx_path)
print(f"Generated DOCX: {docx_path}")
